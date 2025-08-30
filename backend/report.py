from __future__ import annotations
import os
from docx import Document
from docx.shared import Inches
from .schemas import AnalysisResult

def build_report_docx(result: AnalysisResult, out_dir: str) -> str:
    os.makedirs(out_dir, exist_ok=True)
    base = os.path.splitext(os.path.basename(result.filename))[0]
    out_path = os.path.join(out_dir, f"{base}_CIELレポート.docx")
    doc = Document()
    doc.add_heading("👑 CIEL Contract 自動解析レポート（Phase 5 完全版）", level=0)
    doc.add_paragraph(f"ファイル名: {result.filename}")
    doc.add_paragraph(result.summary)
    # Entities
    doc.add_heading("当事者・準拠法・管轄", level=1)
    t = doc.add_table(rows=6, cols=2)
    t.cell(0,0).text = "甲"; t.cell(0,1).text = result.entities.parties.get("甲") or "(未検出)"
    t.cell(1,0).text = "乙"; t.cell(1,1).text = result.entities.parties.get("乙") or "(未検出)"
    t.cell(2,0).text = "丙"; t.cell(2,1).text = result.entities.parties.get("丙") or "(未検出)"
    t.cell(3,0).text = "丁"; t.cell(3,1).text = result.entities.parties.get("丁") or "(未検出)"
    t.cell(4,0).text = "準拠法"; t.cell(4,1).text = result.entities.governing_law or "(未検出)"
    t.cell(5,0).text = "管轄裁判所"; t.cell(5,1).text = result.entities.jurisdiction or "(未検出)"
    # Liability cap & unilateral termination
    doc.add_heading("責任上限・片務解除", level=1)
    cap_txt = f"{result.liability_cap_yen:,} 円" if result.liability_cap_yen is not None else "（未検出）"
    doc.add_paragraph(f"責任上限（数値判定）: {cap_txt}")
    doc.add_paragraph(f"片務解除: {'有' if result.unilateral_termination else '無'}")
    if result.unilateral_clause_numbers:
        doc.add_paragraph(f"該当条文: {', '.join(map(str, result.unilateral_clause_numbers))}")
    # Visuals
    doc.add_heading("リスクヒートマップ", level=1)
    if result.heatmap_path and os.path.exists(result.heatmap_path):
        doc.add_picture(result.heatmap_path, width=Inches(2.2))
    else:
        doc.add_paragraph("(該当カテゴリなし)")
    doc.add_heading("因果マップ（簡易）", level=1)
    if result.causal_graph_path and os.path.exists(result.causal_graph_path):
        doc.add_picture(result.causal_graph_path, width=Inches(5.5))
    else:
        doc.add_paragraph("(未生成)")
    # Clauses
    doc.add_heading("条文ごとの分類・リスク", level=1)
    for c in result.clauses:
        doc.add_heading(f"[{c.index}] {c.category}（信頼度:{c.confidence:.2f} / リスク:{c.risk:.1f}）", level=2)
        doc.add_paragraph(c.text)
    doc.add_page_break()
    doc.add_paragraph("※ 本レポートはローカル自動解析の結果であり、法的助言を構成しません。")
    doc.save(out_path)
    return out_path
